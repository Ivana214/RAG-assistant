import io
from pathlib import Path

import pandas as pd
from docx import Document as DocxDocument
from pptx import Presentation

from app.services.pdf import PDFProcessingError, extract_pdf_pages


class ParserError(ValueError):
    pass


SUPPORTED_EXTENSIONS = {".pdf", ".xlsx", ".csv", ".docx", ".pptx"}


def parse_file(filename: str, content: bytes) -> dict:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ParserError(f"不支持的文件类型: {suffix}")

    if suffix == ".pdf":
        return _parse_pdf(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content)
    if suffix == ".csv":
        return _parse_csv(content)
    if suffix == ".docx":
        return _parse_docx(content)
    return _parse_pptx(content)


def _parse_pdf(content: bytes) -> dict:
    try:
        result = extract_pdf_pages(content)
    except PDFProcessingError as exc:
        raise ParserError(str(exc)) from exc

    items = [
        {
            "text": page["text"],
            "page_number": page["page_number"],
            "metadata": {"format": "pdf", "page": page["page_number"]},
        }
        for page in result["pages"]
    ]
    return {
        "format": "pdf",
        "total_items": len(items),
        "items": items,
        "summary": {"pages": result["total_pages"], "text_pages": len(items)},
    }


def _dataframe_text(label: str, dataframe: pd.DataFrame) -> str:
    header = f"表格: {label}\n列: {', '.join(str(column) for column in dataframe.columns)}\n"
    return header + dataframe.to_string(index=False)


def _missing_summary(dataframe: pd.DataFrame) -> dict:
    return {
        str(column): int(dataframe[column].isna().sum())
        for column in dataframe.columns
        if int(dataframe[column].isna().sum()) > 0
    }


def _parse_xlsx(content: bytes) -> dict:
    try:
        sheets = pd.read_excel(io.BytesIO(content), sheet_name=None)
    except Exception as exc:
        raise ParserError("无法解析 Excel 文件") from exc

    items = []
    summary = []
    for sheet_name, dataframe in sheets.items():
        columns = [str(column) for column in dataframe.columns]
        missing = _missing_summary(dataframe)
        items.append(
            {
                "text": _dataframe_text(sheet_name, dataframe),
                "page_number": len(items) + 1,
                "metadata": {
                    "format": "xlsx",
                    "sheet": sheet_name,
                    "rows": int(dataframe.shape[0]),
                    "columns": columns,
                    "missing": missing,
                },
            }
        )
        summary.append(
            {
                "sheet": sheet_name,
                "rows": int(dataframe.shape[0]),
                "columns": columns,
                "missing": missing,
            }
        )

    return {
        "format": "xlsx",
        "total_items": len(items),
        "items": items,
        "summary": summary,
    }


def _read_csv(content: bytes) -> pd.DataFrame:
    try:
        return pd.read_csv(io.BytesIO(content), encoding="utf-8")
    except UnicodeDecodeError:
        return pd.read_csv(io.BytesIO(content), encoding="gbk")


def _parse_csv(content: bytes) -> dict:
    try:
        dataframe = _read_csv(content)
    except Exception as exc:
        raise ParserError("无法解析 CSV 文件") from exc

    columns = [str(column) for column in dataframe.columns]
    missing = _missing_summary(dataframe)
    item = {
        "text": _dataframe_text("csv", dataframe),
        "page_number": 1,
        "metadata": {
            "format": "csv",
            "rows": int(dataframe.shape[0]),
            "columns": columns,
            "missing": missing,
        },
    }
    return {
        "format": "csv",
        "total_items": 1,
        "items": [item],
        "summary": {
            "rows": int(dataframe.shape[0]),
            "columns": columns,
            "missing": missing,
        },
    }


def _parse_docx(content: bytes) -> dict:
    try:
        document = DocxDocument(io.BytesIO(content))
    except Exception as exc:
        raise ParserError("无法解析 Word 文件") from exc

    items = []
    paragraph_text = "\n".join(
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )
    if paragraph_text:
        items.append(
            {
                "text": paragraph_text,
                "page_number": 1,
                "metadata": {"format": "docx", "kind": "paragraphs"},
            }
        )

    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [
                cell.text.strip().replace("\n", " ") for cell in row.cells
            ]
            rows.append(" | ".join(cells))
        text = "\n".join(rows).strip()
        if text:
            items.append(
                {
                    "text": text,
                    "page_number": len(items) + 1,
                    "metadata": {
                        "format": "docx",
                        "kind": "table",
                        "table_index": table_index,
                    },
                }
            )

    return {
        "format": "docx",
        "total_items": len(items),
        "items": items,
        "summary": {
            "paragraph_blocks": 1 if paragraph_text else 0,
            "tables": len(document.tables),
        },
    }


def _parse_pptx(content: bytes) -> dict:
    try:
        presentation = Presentation(io.BytesIO(content))
    except Exception as exc:
        raise ParserError("无法解析 PPT 文件") from exc

    items = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs).strip()
                if text:
                    parts.append(text)
        text = "\n".join(parts).strip()
        if text:
            items.append(
                {
                    "text": text,
                    "page_number": slide_index,
                    "metadata": {"format": "pptx", "slide": slide_index},
                }
            )

    return {
        "format": "pptx",
        "total_items": len(items),
        "items": items,
        "summary": {
            "slides": len(presentation.slides),
            "text_slides": len(items),
        },
    }
